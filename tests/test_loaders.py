from __future__ import annotations

from pathlib import Path

import docx
import pytest
from pypdf import PdfWriter

from kbagent.loaders import load_sections


def test_plain_text_and_markdown(tmp_path: Path) -> None:
    path = tmp_path / "a.md"
    path.write_text("# 标题\n\n正文内容", encoding="utf-8")
    sections = load_sections(path)
    assert len(sections) == 1
    assert "正文内容" in sections[0].text
    assert sections[0].page is None


def test_docx_paragraphs_and_tables(tmp_path: Path) -> None:
    document = docx.Document()
    document.add_heading("项目说明", level=1)
    document.add_paragraph("这是第一段。")
    table = document.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "字段"
    table.rows[0].cells[1].text = "值"
    path = tmp_path / "a.docx"
    document.save(path)

    text = "\n".join(section.text for section in load_sections(path))
    assert "这是第一段。" in text
    assert "字段 | 值" in text


def test_blank_pdf_raises(tmp_path: Path) -> None:
    path = tmp_path / "blank.pdf"
    writer = PdfWriter()
    writer.add_blank_page(width=200, height=200)
    with path.open("wb") as handle:
        writer.write(handle)

    with pytest.raises(ValueError, match="没有提取到文本"):
        load_sections(path)


def test_empty_file_raises(tmp_path: Path) -> None:
    path = tmp_path / "empty.txt"
    path.write_text("   \n\n", encoding="utf-8")
    with pytest.raises(ValueError, match="没有提取到文本"):
        load_sections(path)


def test_unsupported_suffix(tmp_path: Path) -> None:
    path = tmp_path / "a.csv"
    path.write_text("a,b", encoding="utf-8")
    with pytest.raises(ValueError, match="不支持"):
        load_sections(path)
