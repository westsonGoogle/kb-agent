"""文档加载:把 md/txt/pdf/docx 统一成带页码的 Section 列表。"""

from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path


@dataclass(slots=True)
class Section:
    text: str
    page: int | None = None


def load_sections(path: Path) -> list[Section]:
    suffix = path.suffix.lower()
    if suffix in {".md", ".markdown", ".txt"}:
        sections = [Section(path.read_text(encoding="utf-8", errors="replace"))]
    elif suffix == ".pdf":
        sections = _load_pdf(path)
    elif suffix == ".docx":
        sections = _load_docx(path)
    else:
        raise ValueError(f"不支持的文件类型: {suffix}")

    sections = [section for section in sections if section.text.strip()]
    if not sections:
        raise ValueError("没有提取到文本(扫描版 PDF 需要先 OCR)")
    return sections


def _load_pdf(path: Path) -> list[Section]:
    from pypdf import PdfReader

    try:
        reader = PdfReader(str(path))
        if reader.is_encrypted:
            reader.decrypt("")
    except Exception as exc:  # noqa: BLE001 - 三方库异常类型不稳定
        raise ValueError(f"PDF 解析失败: {exc}") from exc

    sections = []
    for index, page in enumerate(reader.pages, start=1):
        try:
            text = page.extract_text() or ""
        except Exception:  # noqa: BLE001 - 单页失败不应拖垮整份文档
            continue
        sections.append(Section(text=text, page=index))
    return sections


def _load_docx(path: Path) -> list[Section]:
    import docx

    try:
        document = docx.Document(str(path))
    except Exception as exc:  # noqa: BLE001
        raise ValueError(f"DOCX 解析失败: {exc}") from exc

    parts = [paragraph.text for paragraph in document.paragraphs if paragraph.text.strip()]
    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells]
            if any(cells):
                parts.append(" | ".join(cells))
    return [Section(text="\n".join(parts))]
